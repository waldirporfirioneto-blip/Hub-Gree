import os
import re
import gc
import tempfile
from datetime import datetime as dt
from utils.config_app import CAMINHO_LOGO_COR

class ConsolidadorModel:
    
    def processar_vendas(self, entrada, data_doc, solic, fila_ui):
        try:
            import pypdf
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.drawing.image import Image as ExcelImage
            import math

            pdf_files = []
            if os.path.isdir(entrada):
                for f in os.listdir(entrada):
                    if f.lower().endswith(".pdf") and not f.startswith("~$") and not f.startswith("Mesclado_"):
                        pdf_files.append(os.path.join(entrada, f))
            elif os.path.isfile(entrada) and entrada.lower().endswith(".pdf"):
                pdf_files.append(entrada)

            if not pdf_files:
                fila_ui.put(("error", "Nenhum arquivo PDF válido localizado!"))
                return

            pdf_files.sort()
            fila_ui.put(("log", f"Localizado(s) {len(pdf_files)} arquivo(s) PDF para leitura."))
            total_arquivos = len(pdf_files)
            regex_unificado = re.compile(r'\b([S5]\d{8}|100\d{6}|TPZ\d{6}|TPR\d{6})\b', re.IGNORECASE)
            orders_list = []

            for idx, filepath in enumerate(pdf_files):
                filename = os.path.basename(filepath)
                fila_ui.put(("log", f"Lendo ({idx+1}/{total_arquivos}): {filename}"))
                fila_ui.put(("progress", int(((idx + 1) / total_arquivos) * 90)))
                texto_extraido = ""
                try:
                    with open(filepath, "rb") as f:
                        reader = pypdf.PdfReader(f, strict=False)
                        for page in reader.pages:
                            text = page.extract_text()
                            if text:
                                texto_extraido += text + "\n"
                except Exception:
                    fila_ui.put(("log", f"⚠️ Erro ao ler conteúdo interno de {filename}. Usando fallback..."))
                
                found_any = False
                if texto_extraido:
                    matches = regex_unificado.findall(texto_extraido)
                    if matches:
                        for match in matches:
                            codigo = match.upper()
                            if codigo not in orders_list:
                                orders_list.append(codigo)
                        found_any = True
                
                if not found_any:
                    nome_sem_ext = os.path.splitext(filename)[0].upper()
                    matches = regex_unificado.findall(nome_sem_ext)
                    if matches:
                        fila_ui.put(("log", f"💡 {filename} lido via nome do arquivo (fallback)."))
                        for match in matches:
                            codigo = match.upper()
                            if codigo not in orders_list:
                                orders_list.append(codigo)
                    else:
                        fila_ui.put(("log", f"⚠️ Nenhum padrão encontrado em {filename}."))
                gc.collect()

            if not orders_list:
                fila_ui.put(("error", "Nenhum código encontrado nos arquivos!"))
                return

            fila_ui.put(("log", f"Total de {len(orders_list)} ordens capturadas! Organizando planilha..."))
            orders_ordenados = orders_list

            total_pedidos = len(orders_ordenados)
            itens_col = math.ceil(total_pedidos / 6)
            if itens_col == 0: itens_col = 1

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Controle de Entrega"
            ws.views.sheetView[0].showGridLines = True

            font_title = Font(name="Segoe UI", size=16, bold=True, color="1F4E78")
            font_subtitle = Font(name="Segoe UI", size=12, bold=True, color="595959")
            font_meta = Font(name="Segoe UI", size=10, bold=True, color="000000")
            font_meta_center = Font(name="Segoe UI", size=10, bold=True, color="595959")
            font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            font_data_bold = Font(name="Segoe UI", size=10, bold=True, color="000000")
            font_data_reg = Font(name="Segoe UI", size=10, bold=False, color="000000")
            font_footer_label = Font(name="Segoe UI", size=10, bold=True, color="2F5496")
            font_footer_line = Font(name="Segoe UI", size=11, bold=False, color="595959")
            fill_header = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
            fill_zebra_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
            fill_zebra_gray = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            fill_special = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            thin_side = Side(border_style="thin", color="D9D9D9")
            border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

            ws.column_dimensions['A'].width = 3
            for col_idx in range(6):
                ped_col = chr(ord('B') + col_idx * 2)
                ok_col = chr(ord('C') + col_idx * 2)
                ws.column_dimensions[ped_col].width = 14
                ws.column_dimensions[ok_col].width = 5

            ws.merge_cells("B2:M2")
            ws.row_dimensions[2].height = 45

            try:
                img_logo = ExcelImage(CAMINHO_LOGO_COR)
                ws.add_image(img_logo, "B2")
            except Exception as e:
                fila_ui.put(("log", f"Aviso sobre a logo: {str(e)}"))
                ws["B2"] = "GREE ELECTRIC APPLIANCES DO BRASIL LTDA."
                ws["B2"].font = font_title
                ws["B2"].alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells("B4:M4")
            ws["B4"] = "Controle de Entrega - Pedidos de Vendas"
            ws["B4"].font = font_subtitle
            ws["B4"].alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells("B7:D7")
            ws["B7"] = f"Data do Documento: {data_doc}"
            ws["B7"].font = font_meta
            ws["B7"].alignment = Alignment(horizontal="left", vertical="center")

            ws.merge_cells("G7:I7")
            ws["G7"] = "SISTEMA GREE APP"
            ws["G7"].font = font_meta_center
            ws["G7"].alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells("K7:M7")
            ws["K7"] = f"Solicitação: {solic}"
            ws["K7"].font = font_meta
            ws["K7"].alignment = Alignment(horizontal="right", vertical="center")

            for col_idx in range(6):
                ped_col = chr(ord('B') + col_idx * 2)
                ok_col = chr(ord('C') + col_idx * 2)
                ws[f"{ped_col}9"] = "Pedido"
                ws[f"{ped_col}9"].font = font_header
                ws[f"{ped_col}9"].fill = fill_header
                ws[f"{ped_col}9"].alignment = Alignment(horizontal="center", vertical="center")
                ws[f"{ped_col}9"].border = border_cell
                ws[f"{ok_col}9"] = "OK"
                ws[f"{ok_col}9"].font = font_header
                ws[f"{ok_col}9"].fill = fill_header
                ws[f"{ok_col}9"].alignment = Alignment(horizontal="center", vertical="center")
                ws[f"{ok_col}9"].border = border_cell

            total_slots = itens_col * 6
            for i in range(total_slots):
                col_num = i // itens_col
                row_num = 10 + (i % itens_col)
                ped_col = chr(ord('B') + col_num * 2)
                ok_col = chr(ord('C') + col_num * 2)

                order_val = orders_ordenados[i] if i < len(orders_ordenados) else ""
                is_special = order_val.startswith("100") or order_val.startswith("TPZ") or order_val.startswith("TPR")
                row_fill = fill_zebra_white if (row_num % 2 == 0) else fill_zebra_gray

                ws[f"{ped_col}{row_num}"] = order_val
                ws[f"{ok_col}{row_num}"] = "☐"
                
                if order_val and is_special:
                    ws[f"{ped_col}{row_num}"].font = font_data_bold
                    ws[f"{ped_col}{row_num}"].fill = fill_special
                    ws[f"{ok_col}{row_num}"].font = font_data_bold
                    ws[f"{ok_col}{row_num}"].fill = fill_special
                else:
                    ws[f"{ped_col}{row_num}"].font = font_data_reg
                    ws[f"{ped_col}{row_num}"].fill = row_fill
                    ws[f"{ok_col}{row_num}"].font = font_data_reg
                    ws[f"{ok_col}{row_num}"].fill = row_fill

                ws[f"{ped_col}{row_num}"].alignment = Alignment(horizontal="left", vertical="center")
                ws[f"{ped_col}{row_num}"].border = border_cell
                ws[f"{ok_col}{row_num}"].alignment = Alignment(horizontal="center", vertical="center")
                ws[f"{ok_col}{row_num}"].border = border_cell

            footer_row = 10 + itens_col + 2
            ws.merge_cells(f"B{footer_row}:E{footer_row}")
            ws[f"B{footer_row}"] = "________________________________________________"
            ws[f"B{footer_row}"].font = font_footer_line
            ws[f"B{footer_row}"].alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells(f"I{footer_row}:L{footer_row}")
            ws[f"I{footer_row}"] = "______ / ______ / 2026"
            ws[f"I{footer_row}"].font = font_footer_line
            ws[f"I{footer_row}"].alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells(f"B{footer_row+1}:E{footer_row+1}")
            ws[f"B{footer_row+1}"] = "Responsável"
            ws[f"B{footer_row+1}"].font = font_footer_label
            ws[f"B{footer_row+1}"].alignment = Alignment(horizontal="center", vertical="center")

            ws.merge_cells(f"I{footer_row+1}:L{footer_row+1}")
            ws[f"I{footer_row+1}"] = "Data de Recebimento"
            ws[f"I{footer_row+1}"].font = font_footer_label
            ws[f"I{footer_row+1}"].alignment = Alignment(horizontal="center", vertical="center")

            diretorio_destino = entrada if os.path.isdir(entrada) else os.path.dirname(entrada)
            filename_output = f"controle_de_entrega_venda_Gree_{dt.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            out_xlsx = os.path.join(diretorio_destino, filename_output)

            wb.save(out_xlsx)
            wb.close()
            gc.collect()

            fila_ui.put(("progress", 100))
            fila_ui.put(("success", f"Relatório gerado com sucesso!\nSalvo em: {out_xlsx}"))

        except Exception as e:
            fila_ui.put(("error", f"Ocorreu um erro no processamento: {str(e)}"))

    def mesclar_pdfs(self, arquivos_lista, dest_path, fila_ui):
        try:
            import pypdf
            writer = pypdf.PdfWriter()
            total_arquivos = len(arquivos_lista)
            arquivos_anexados = 0
            
            for idx, file_path in enumerate(arquivos_lista):
                filename = os.path.basename(file_path)
                if os.path.normpath(file_path) == os.path.normpath(dest_path): continue
                
                fila_ui.put(("log", f"Anexando PDF ({idx+1}/{total_arquivos}): {filename}"))
                fila_ui.put(("progress", int(((idx + 1) / total_arquivos) * 90)))
                
                try:
                    with open(file_path, "rb") as f_in:
                        reader = pypdf.PdfReader(f_in, strict=False)
                        for page in reader.pages:
                            writer.add_page(page)
                        arquivos_anexados += 1
                except Exception as ex:
                    fila_ui.put(("log", f"⚠️ Arquivo ignorado: {filename} ({str(ex)})"))
                    
            if arquivos_anexados == 0:
                fila_ui.put(("error", "Nenhum arquivo válido pôde ser lido!"))
                return
                
            fila_ui.put(("log", "Gravando arquivo unificado..."))
            with open(dest_path, "wb") as f_out:
                writer.write(f_out)
            writer.close()
            gc.collect()
            
            fila_ui.put(("progress", 100))
            fila_ui.put(("success", f"Mesclagem finalizada!\nArquivo gravado: {dest_path}"))
        except Exception as e:
            fila_ui.put(("error", f"Erro ao mesclar: {str(e)}"))

    def converter_pdf_para_word(self, pdf_path, fila_ui):
        try:
            import pypdf
            import docx
            from docx import Document
            
            pypdfium_available = False
            try:
                import pypdfium2 as pdfium
                pypdfium_available = True
            except ImportError: pass

            fila_ui.put(("log", f"Iniciando conversão de PDF para Word..."))
            filename = os.path.basename(pdf_path)
            dir_name = os.path.dirname(pdf_path)
            base_name = os.path.splitext(filename)[0]
            out_docx = os.path.join(dir_name, f"{base_name}.docx")

            try:
                if os.path.exists(out_docx):
                    with open(out_docx, "a+b") as f: pass
            except IOError:
                fila_ui.put(("error", "O arquivo Word de destino já está aberto!"))
                return

            doc = Document()
            doc.add_heading(base_name, level=0)

            with open(pdf_path, "rb") as f:
                reader = pypdf.PdfReader(f, strict=False)
                total_pages = len(reader.pages)
                has_any_text = False
                for page in reader.pages:
                    t = page.extract_text()
                    if t and t.strip():
                        has_any_text = True
                        break

                if has_any_text:
                    fila_ui.put(("log", "Extraindo parágrafos..."))
                    for i, page in enumerate(reader.pages):
                        fila_ui.put(("log", f"Lendo página ({i+1}/{total_pages})..."))
                        fila_ui.put(("progress", int(((i + 1) / total_pages) * 90)))
                        text = page.extract_text()
                        if text:
                            for paragraph in text.split("\n"):
                                if paragraph.strip():
                                    doc.add_paragraph(paragraph)
                        if i < total_pages - 1:
                            doc.add_page_break()
                else:
                    if pypdfium_available:
                        fila_ui.put(("log", "Convertendo páginas escaneadas..."))
                        pdf_doc = pdfium.PdfDocument(pdf_path)
                        for i in range(total_pages):
                            fila_ui.put(("log", f"Processando ({i+1}/{total_pages})..."))
                            fila_ui.put(("progress", int(((i + 1) / total_pages) * 90)))
                            page = pdf_doc[i]
                            bitmap = page.render(scale=2.0)
                            pil_img = bitmap.to_pil()
                            with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as tmp_file:
                                tmp_path = tmp_file.name
                            pil_img.save(tmp_path, "JPEG")
                            doc.add_picture(tmp_path, width=docx.shared.Inches(5.8))
                            try: os.remove(tmp_path)
                            except: pass
                            if i < total_pages - 1: doc.add_page_break()
                        pdf_doc.close()
                    else:
                        doc.add_paragraph("Este é um PDF escaneado (imagem). pypdfium2 necessário para imagens.")

            fila_ui.put(("log", "Gravando documento..."))
            doc.save(out_docx)
            fila_ui.put(("progress", 100))
            fila_ui.put(("success", f"PDF convertido com sucesso!\nSalvo em: {out_docx}"))
        except Exception as e:
            fila_ui.put(("error", f"Erro na conversão: {str(e)}"))

    def converter_pdf_para_jpg(self, pdf_path, dpi, fila_ui):
        try:
            import pypdfium2 as pdfium
            fila_ui.put(("log", f"Extraindo imagens via pypdfium2..."))
            filename = os.path.basename(pdf_path)
            dir_name = os.path.dirname(pdf_path)
            base_name = os.path.splitext(filename)[0]
            out_dir = os.path.join(dir_name, f"Imagens_{base_name}")
            os.makedirs(out_dir, exist_ok=True)

            doc = pdfium.PdfDocument(pdf_path)
            total_pages = len(doc)
            zoom = dpi / 72.0

            for i in range(total_pages):
                fila_ui.put(("log", f"Renderizando {i+1}/{total_pages}..."))
                fila_ui.put(("progress", int(((i + 1) / total_pages) * 90)))
                page = doc[i]
                bitmap = page.render(scale=zoom)
                pil_img = bitmap.to_pil()
                out_file = os.path.join(out_dir, f"pagina_{i+1:03d}.jpg")
                pil_img.save(out_file, "JPEG")

            doc.close()
            fila_ui.put(("progress", 100))
            fila_ui.put(("success", f"Convertido com sucesso!\nPasta: {out_dir}"))
        except Exception as e:
            fila_ui.put(("error", f"Erro: {str(e)}"))

    def converter_jpg_para_pdf(self, images_list, output_path, orientation, margin, fila_ui):
        try:
            from PIL import Image
            fila_ui.put(("log", f"Iniciando compilação..."))
            total_imgs = len(images_list)
            pil_images = []

            for i, img_path in enumerate(images_list):
                fila_ui.put(("log", f"Processando ({i+1}/{total_imgs})..."))
                fila_ui.put(("progress", int(((i + 1) / total_imgs) * 80)))
                img = Image.open(img_path)

                if img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info):
                    fundo_branco = Image.new("RGB", img.size, "white")
                    if img.mode == "P": img = img.convert("RGBA")
                    fundo_branco.paste(img, mask=img.split()[-1])
                    img = fundo_branco
                elif img.mode != "RGB":
                    img = img.convert("RGB")

                if orientation == "Paisagem (Horizontal)":
                    a4_w, a4_h = 3508, 2480
                else:
                    a4_w, a4_h = 2480, 3508

                if margin == "Margem fina": margin_px = 80
                elif margin == "Margem larga": margin_px = 250
                else: margin_px = 0

                target_w, target_h = a4_w - (2 * margin_px), a4_h - (2 * margin_px)
                img_w, img_h = img.size
                ratio = min(target_w / img_w, target_h / img_h)
                new_w, new_h = int(img_w * ratio), int(img_h * ratio)

                resized_img = img.resize((new_w, new_h), Image.Resampling.LANCZOS)
                canvas = Image.new("RGB", (a4_w, a4_h), "white")
                offset_x, offset_y = (a4_w - new_w) // 2, (a4_h - new_h) // 2
                canvas.paste(resized_img, (offset_x, offset_y))
                pil_images.append(canvas)

            if not pil_images:
                fila_ui.put(("error", "Nenhuma imagem lida!"))
                return

            fila_ui.put(("log", "Salvando PDF em HD (300 DPI)..."))
            fila_ui.put(("progress", 90))

            pil_images[0].save(output_path, "PDF", resolution=300.0, save_all=True, append_images=pil_images[1:], quality=100, subsampling=0)
            fila_ui.put(("progress", 100))
            fila_ui.put(("success", f"PDF gerado em Alta Definição!\nSalvo em: {output_path}"))
        except Exception as e:
            fila_ui.put(("error", f"Erro: {str(e)}"))