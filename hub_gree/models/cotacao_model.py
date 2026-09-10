import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter

# Funções auxiliares isoladas no modelo
def letter_to_index(letter):
    letter = letter.upper()
    result = 0
    for char in letter:
        result = result * 26 + (ord(char) - ord('A') + 1)
    return result - 1

def safe_str(val):
    if pd.isna(val):
        return ''
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()

class CotacaoModel:
    def __init__(self):
        self.df_dados = None
        self.df_original_cif = None
        self.is_gmax = False

    def processar_base_dados(self, caminho_entrada):
        xls = pd.ExcelFile(caminho_entrada)
        if "Programação Filtrada" not in xls.sheet_names:
            raise ValueError("A aba 'Programação Filtrada' não foi encontrada no arquivo.")

        df_bruto = pd.read_excel(caminho_entrada, sheet_name="Programação Filtrada")
        col_frete = df_bruto.columns[letter_to_index('G')]
        df_cif = df_bruto[df_bruto[col_frete].astype(str).str.upper().str.strip() == 'CIF'].copy()

        if df_cif.empty:
            raise ValueError("Nenhuma carga 'CIF' encontrada na base.")

        self.df_original_cif = df_cif.copy()
        self.is_gmax = False

        if len(df_bruto.columns) > 32:
            col_nome_ag = str(df_bruto.columns[letter_to_index('AG')]).strip().lower()
            if 'venda' in col_nome_ag or 'cod' in col_nome_ag:
                self.is_gmax = True

        visao_texto = "G-MAX" if self.is_gmax else "PADRÃO"
        col_tipo_veiculo = df_cif.columns[letter_to_index('F')]
        col_item = df_cif.columns[letter_to_index('A')]
        
        df_cif[col_tipo_veiculo] = df_cif.groupby(col_item)[col_tipo_veiculo].transform(lambda x: x.ffill().bfill())
        col_y_idx = letter_to_index('Y')
        df_cif['CICLO_calc'] = df_cif.iloc[:, col_y_idx].fillna('') if col_y_idx < len(df_bruto.columns) else ''

        if self.is_gmax:
            mapa = {'ITEM': 'A', 'CIDADE': 'R', 'UF': 'S', 'CLIENTE': 'P', 'CODIGO': 'AG', 'QTD': 'AA', 'M3': 'AD', 'VALOR': 'AF', 'DESCRICAO': 'AK'}
        else:
            mapa = {'ITEM': 'A', 'CIDADE': 'R', 'UF': 'S', 'CLIENTE': 'P', 'CODIGO': 'AH', 'QTD': 'AB', 'M3': 'AE', 'VALOR': 'AG', 'DESCRICAO': 'AJ'}

        df_cif['ITEM_calc'] = df_cif.iloc[:, letter_to_index(mapa['ITEM'])].apply(safe_str)
        df_cif['CIDADE_calc'] = df_cif.iloc[:, letter_to_index(mapa['CIDADE'])].fillna('')
        df_cif['UF_calc'] = df_cif.iloc[:, letter_to_index(mapa['UF'])].fillna('')
        df_cif['CLIENTE_calc'] = df_cif.iloc[:, letter_to_index(mapa['CLIENTE'])].fillna('')
        df_cif['CODIGO_calc'] = df_cif.iloc[:, letter_to_index(mapa['CODIGO'])].fillna('')
        df_cif['QTD_calc'] = df_cif.iloc[:, letter_to_index(mapa['QTD'])].fillna(0)
        df_cif['M3_calc'] = df_cif.iloc[:, letter_to_index(mapa['M3'])].fillna(0.0)
        df_cif['VALOR_calc'] = df_cif.iloc[:, letter_to_index(mapa['VALOR'])].fillna(0.0)
        df_cif['DESCRICAO_calc'] = df_cif.iloc[:, letter_to_index(mapa['DESCRICAO'])].fillna('')
        df_cif['TIPO_VEICULO_calc'] = df_cif[col_tipo_veiculo].fillna('')

        self.df_dados = df_cif
        itens_unicos = [str(item) for item in df_cif['ITEM_calc'].drop_duplicates().tolist() if str(item).strip() != ""]

        return itens_unicos, visao_texto, len(df_cif)

    def gerar_planilha(self, arquivo_saida, ordem_desejada):
        wb = Workbook()
        wb.remove(wb.active)

        mascara_container = self.df_dados['TIPO_VEICULO_calc'].astype(str).str.contains('Container', case=False, na=False)
        df_cabotagem = self.df_dados[mascara_container]
        df_rodoviario = self.df_dados[~mascara_container]

        if not df_rodoviario.empty:
            ws_rodo = wb.create_sheet(title="Cotação Rodoviário")
            self._gerar_aba_excel(ws_rodo, df_rodoviario, "RODOVIARIO", ordem_desejada)

        if not df_cabotagem.empty:
            ws_cabo = wb.create_sheet(title="Cotação Cabotagem")
            self._gerar_aba_excel(ws_cabo, df_cabotagem, "CABOTAGEM", ordem_desejada)

        ws_cif_copy = wb.create_sheet(title="Programação Filtrada (CIF)")
        df_export = self.df_original_cif.copy()
        
        for col in df_export.columns:
            if str(df_export[col].dtype).startswith('datetime'):
                df_export[col] = df_export[col].dt.strftime('%d/%m/%Y').fillna('')
            else:
                df_export[col] = df_export[col].fillna('')

        for r in dataframe_to_rows(df_export, index=False, header=True):
            ws_cif_copy.append(r)

        wb.save(arquivo_saida)

    def _gerar_aba_excel(self, ws, df_filtrado, tipo_aba, ordem_desejada):
        # A exata lógica visual do Excel que você criou
        fonte_branca_normal = Font(bold=False, color="FFFFFF")
        fonte_preta_normal = Font(bold=False, color="000000")
        fonte_preta_negrito = Font(bold=True, color="000000")
        fill_azul_escuro = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
        fill_cinza = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
        fill_amarelo = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        alinhamento_centro = Alignment(horizontal="center", vertical="center")
        borda_fina = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        if tipo_aba == "CABOTAGEM":
            transp = ["COSTA BRASIL", "CENTER CARGO", "TECMAR"]
        else:
            transp = ["TRANSBUIATTE", "SPEED", "SR LOG", "BUSSOLA", "TODOBRASIL", "AMAZON", "GAB", "BERTOLINI"]

        titulos = ["ITEM"] + transp + ["CIDADE", "UF", "CLIENTE", "CODIGO", "QTD", "M³", "Valor", "DESCRIÇÃO"]
        titulos_colunas = {i+1: nome for i, nome in enumerate(titulos)}
        num_transp = len(transp)

        col_cidade = titulos.index("CIDADE") + 1
        col_uf = titulos.index("UF") + 1
        col_cliente = titulos.index("CLIENTE") + 1
        col_codigo = titulos.index("CODIGO") + 1
        col_qtd = titulos.index("QTD") + 1
        col_m3 = titulos.index("M³") + 1
        col_valor = titulos.index("Valor") + 1
        col_descricao = titulos.index("DESCRIÇÃO") + 1

        linha_atual = 1

        for item_nome in ordem_desejada:
            grupo = df_filtrado[df_filtrado['ITEM_calc'] == item_nome]
            if grupo.empty: continue

            for col_idx, col_nome in titulos_colunas.items():
                celula = ws.cell(row=linha_atual, column=col_idx, value=col_nome)
                celula.fill = fill_azul_escuro
                celula.font = fonte_branca_normal
                celula.alignment = alinhamento_centro
                celula.border = borda_fina

            linha_atual += 1
            linha_inicial = linha_atual

            for _, row in grupo.iterrows():
                ws.cell(row=linha_atual, column=col_cidade, value=row.get('CIDADE_calc', ''))
                ws.cell(row=linha_atual, column=col_uf, value=row.get('UF_calc', ''))
                ws.cell(row=linha_atual, column=col_cliente, value=row.get('CLIENTE_calc', ''))
                ws.cell(row=linha_atual, column=col_codigo, value=row.get('CODIGO_calc', ''))

                cel_qtd = ws.cell(row=linha_atual, column=col_qtd, value=row.get('QTD_calc', 0))
                cel_qtd.number_format = '0'

                cel_m3 = ws.cell(row=linha_atual, column=col_m3, value=row.get('M3_calc', 0.0))
                cel_m3.number_format = '0.00'

                is_painel = False
                if not self.is_gmax:
                    ciclo_info = str(row.get('CICLO_calc', '')).strip().upper()
                    if 'PAINEL' in ciclo_info:
                        is_painel = True

                if is_painel:
                    ws.cell(row=linha_atual, column=col_valor, value="")
                else:
                    cel_valor_cel = ws.cell(row=linha_atual, column=col_valor, value=row.get('VALOR_calc', 0.0))
                    cel_valor_cel.number_format = '"R$" #,##0.00'

                ws.cell(row=linha_atual, column=col_descricao, value=row.get('DESCRICAO_calc', ''))

                for c in range(1, len(titulos) + 1):
                    cel = ws.cell(row=linha_atual, column=c)
                    cel.border = borda_fina
                    if c >= col_cidade: cel.alignment = alinhamento_centro

                linha_atual += 1

            linha_final = linha_atual - 1
            linha_total = linha_atual

            ws.cell(row=linha_inicial, column=1, value=item_nome)
            letra_valor = get_column_letter(col_valor)

            for col in range(1, num_transp + 2):
                ws.merge_cells(start_row=linha_inicial, start_column=col, end_row=linha_final, end_column=col)
                cel_mesclada = ws.cell(row=linha_inicial, column=col)

                for r in range(linha_inicial, linha_final + 1):
                    c_borda = ws.cell(row=r, column=col)
                    c_borda.border = borda_fina
                    c_borda.alignment = alinhamento_centro

                if col >= 2:
                    transp_nome = titulos[col - 1]
                    if transp_nome == "TRANSBUIATTE":
                        num_entregas = len(grupo[['CIDADE_calc', 'CLIENTE_calc']].drop_duplicates())
                        acrescimo = max(0, num_entregas - 1) * 1100
                        cidades_bloco = grupo['CIDADE_calc'].astype(str).str.upper().str.strip().tolist()
                        is_manaus = any('MANAUS' in c for c in cidades_bloco)

                        if is_manaus:
                            formula_transb = f"=4500+{acrescimo}" if acrescimo > 0 else "=4500"
                        else:
                            formula_transb = f"=${letra_valor}${linha_total}*0.04+{acrescimo}" if acrescimo > 0 else f"=${letra_valor}${linha_total}*0.04"
                        cel_mesclada.value = formula_transb
                    cel_mesclada.number_format = '"R$" #,##0.00'

            for col in range(1, num_transp + 2):
                cel_tot_esq = ws.cell(row=linha_total, column=col)
                cel_tot_esq.fill = fill_cinza
                cel_tot_esq.font = fonte_preta_normal
                cel_tot_esq.border = borda_fina
                cel_tot_esq.alignment = alinhamento_centro
                if col == 1:
                    cel_tot_esq.value = "% Frete"
                else:
                    letra_col = get_column_letter(col)
                    cel_tot_esq.value = f"={letra_col}{linha_inicial}/${letra_valor}${linha_total}"
                    cel_tot_esq.number_format = '0.00%'

            letra_m3 = get_column_letter(col_m3)
            cel_soma_m3 = ws.cell(row=linha_total, column=col_m3, value=f"=SUM({letra_m3}{linha_inicial}:{letra_m3}{linha_final})")
            cel_soma_m3.fill = fill_amarelo
            cel_soma_m3.font = fonte_preta_negrito
            cel_soma_m3.border = borda_fina
            cel_soma_m3.number_format = '0.00'
            cel_soma_m3.alignment = alinhamento_centro

            cel_soma_valor = ws.cell(row=linha_total, column=col_valor, value=f"=SUM({letra_valor}{linha_inicial}:{letra_valor}{linha_final})")
            cel_soma_valor.fill = fill_amarelo
            cel_soma_valor.font = fonte_preta_negrito
            cel_soma_valor.border = borda_fina
            cel_soma_valor.number_format = '"R$" #,##0.00'
            cel_soma_valor.alignment = alinhamento_centro

            linha_atual = linha_total + 3

        larguras_padrao = {
            "ITEM": 15, "TRANSBUIATTE": 15, "SPEED": 12, "SR LOG": 12, "BUSSOLA": 12,
            "TODOBRASIL": 15, "AMAZON": 12, "GAB": 15, "BERTOLINI": 18,
            "COSTA BRASIL": 18, "CENTER CARGO": 18, "TECMAR": 18,
            "CIDADE": 20, "UF": 5, "CLIENTE": 45, "CODIGO": 15,
            "QTD": 8, "M³": 10, "Valor": 18, "DESCRIÇÃO": 35
        }
        for col_idx, col_nome in titulos_colunas.items():
            letra = get_column_letter(col_idx)
            ws.column_dimensions[letra].width = larguras_padrao.get(col_nome, 15)