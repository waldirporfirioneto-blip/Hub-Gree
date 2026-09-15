import os
import openpyxl
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl.utils import get_column_letter
import datetime
import traceback
import pythoncom
from utils.excel_com_engine import abrir_excel_silencioso, abrir_workbook_com_retry, obter_campo_dinamico, _normalizar_titulo_planilha

LISTA_CORES = ["8EA9DB", "A9D08E", "F4B084", "FF99CC", "B4C6E7", "FFD966", "C6E0B4", "CC99FF", "E2EFDA", "99CCFF", "F8CBAD", "D9E1F2"]

def converter_numero(valor):
    if valor is None or str(valor).strip() == "": return None
    try: return int(str(valor).strip().replace(".", "").replace(",", ""))
    except: return valor

def moeda_americana_para_brasileira(valor):
    if valor is None or valor == "": return None
    if isinstance(valor, (int, float)): return float(valor)
    texto = str(valor).strip().replace(",", "")
    try: return float(texto)
    except: return None

def converter_cubagem(valor):
    if valor is None or valor == "": return None
    if isinstance(valor, (int, float)): return float(valor)
    texto = str(valor).strip().replace(".", ",") if "." in str(valor).strip() and "," not in str(valor).strip() else str(valor).strip()
    texto = texto.replace(".", "").replace(",", ".")
    try: return float(texto)
    except: return None

def remover_sufixo(valor):
    if valor is None: return ""
    texto = str(valor).strip()
    return texto[:-2] if len(texto) >= 2 and (texto.endswith("-1") or texto.endswith("-2")) else texto

class CRMModel:
    def analisar_arquivos(self, arquivos):
        mapeamento_obs = []
        datas = set()
        for caminho in arquivos:
            wb = openpyxl.load_workbook(caminho, data_only=True)
            ws = wb.active
            cabecalhos = [str(ws.cell(row=1, column=c).value).strip().lower() for c in range(1, ws.max_column + 1)]
            idx_cliente = -1
            idx_adic = -1

            for i, cab in enumerate(cabecalhos, 1):
                cab_clean = cab.replace(".", "").replace(" ", "").strip()
                if "cliente" in cab_clean and "obs" in cab_clean: idx_cliente = i
                elif "adicionais" in cab_clean and ("obs" in cab_clean or "observa" in cab_clean): idx_adic = i

            if idx_cliente == -1 or idx_adic == -1:
                for i, cab in enumerate(cabecalhos, 1):
                    cab_clean = cab.replace(".", "").replace(" ", "").strip()
                    if "obs" in cab_clean or "observa" in cab_clean:
                        if i != idx_cliente and i != idx_adic:
                            if idx_cliente == -1: idx_cliente = i
                            elif idx_adic == -1: idx_adic = i

            for i in range(2, ws.max_row + 1):
                v = ws.cell(row=i, column=3).value
                if v:
                    d = v.strftime('%d/%m/%Y') if isinstance(v, (datetime.date, datetime.datetime)) else str(v).strip()
                    if d:
                        datas.add(d)
                        val_cli = ws.cell(row=i, column=idx_cliente).value if idx_cliente != -1 else None
                        txt_cli = str(val_cli).strip() if val_cli is not None else ""
                        val_adi = ws.cell(row=i, column=idx_adic).value if idx_adic != -1 else None
                        txt_adi = str(val_adi).strip() if val_adi is not None else ""
                        
                        mapeamento_obs.append({
                            'data': d,
                            'cliente': txt_cli if txt_cli else "(Vazio)",
                            'adic': txt_adi if txt_adi else "(Vazio)"
                        })
            wb.close()
        return sorted(list(datas)), mapeamento_obs

    def obter_tipo_planilha(self, arquivo_path):
        wb = openpyxl.load_workbook(arquivo_path, data_only=True)
        ws = wb.active
        cabecalhos = [str(ws.cell(row=1, column=c).value).strip() for c in range(1, ws.max_column + 1)]
        wb.close()
        return "G-MAX" if "Código Unidade" in cabecalhos or "Unidade" in cabecalhos else "PADRÃO"

    def formatar_padrao(self, arquivo_path, data_escolhida, caminho_final, obs_cliente, obs_adic, cb_progresso):
        cb_progresso(0.1)
        wb = openpyxl.load_workbook(arquivo_path)
        ws = wb.active
        qtd_colunas_originais = ws.max_column
        cabecalhos = [str(ws.cell(row=1, column=c).value).strip().lower() for c in range(1, ws.max_column + 1)]
        idx_obs_cliente, idx_obs_adc = -1, -1

        for i, cab in enumerate(cabecalhos):
            cab_clean = cab.replace(".", "").replace(" ", "").strip()
            if "cliente" in cab_clean and "obs" in cab_clean: idx_obs_cliente = i
            elif "adicionais" in cab_clean and ("obs" in cab_clean or "observa" in cab_clean): idx_obs_adc = i

        if idx_obs_cliente == -1 or idx_obs_adc == -1:
            for i, cab in enumerate(cabecalhos):
                cab_clean = cab.replace(".", "").replace(" ", "").strip()
                if "obs" in cab_clean or "observa" in cab_clean:
                    if i != idx_obs_cliente and i != idx_obs_adc:
                        if idx_obs_cliente == -1: idx_obs_cliente = i
                        elif idx_obs_adc == -1: idx_obs_adc = i

        ws_copia = wb.copy_worksheet(ws)
        ws_copia.title = "Programação Completo"
        ws.title = "Programação Filtrada"

        linhas_raw = list(ws.iter_rows(min_row=2, values_only=True))
        linhas_validas = []
        req_len = max(ws.max_column, 65)
        ultimo_valor_conhecido = {i: None for i in range(9, 14)}

        cb_progresso(0.2)

        for r_tuple in linhas_raw:
            r = list(r_tuple)
            if len(r) < req_len: r.extend([None] * (req_len - len(r)))

            v_cliente = r[idx_obs_cliente] if idx_obs_cliente != -1 and idx_obs_cliente < len(r) else None
            v_adc = r[idx_obs_adc] if idx_obs_adc != -1 and idx_obs_adc < len(r) else None

            txt_cliente = str(v_cliente).strip() if v_cliente is not None else ""
            txt_adc = str(v_adc).strip() if v_adc is not None else ""
            txt_cliente_check = txt_cliente if txt_cliente else "(Vazio)"
            txt_adc_check = txt_adc if txt_adc else "(Vazio)"

            pular_linha = False
            val_nf = ""
            if obs_cliente or obs_adic:
                match_cliente = (txt_cliente_check in obs_cliente) if obs_cliente else True
                match_adic = (txt_adc_check in obs_adic) if obs_adic else True
                if not (match_cliente and match_adic):
                    pular_linha = True
                else:
                    if "nf" in txt_cliente.lower(): val_nf = txt_cliente_check
                    elif "nf" in txt_adc.lower(): val_nf = txt_adc_check
                    else:
                        val_nf = txt_cliente_check if txt_cliente_check != "(Vazio)" else txt_adc_check
                        if val_nf == "(Vazio)": val_nf = ""
            else:
                if "nf" in txt_cliente.lower(): val_nf = txt_cliente
                elif "nf" in txt_adc.lower(): val_nf = txt_adc
                else: val_nf = txt_cliente if txt_cliente else txt_adc

            if pular_linha: continue

            for i in range(9, 14):
                if r[i] is None or str(r[i]).strip() == "": r[i] = ultimo_valor_conhecido[i]
                else: ultimo_valor_conhecido[i] = r[i]

            r[0] = converter_numero(r[0])
            for i in [26, 32, 39, 42]: r[i] = moeda_americana_para_brasileira(r[i])
            r[30] = converter_cubagem(r[30])
            r[17] = r[50]
            r[18] = r[51]
            val_al = r[37]
            r[37] = remover_sufixo(val_al) if val_al and str(val_al).strip() != "" else r[33]
            r[40] = remover_sufixo(r[40])

            v = r[2]
            d = v.strftime('%d/%m/%Y') if isinstance(v, (datetime.date, datetime.datetime)) else str(v).strip()
            if d == data_escolhida:
                try: r[2] = datetime.datetime.strptime(data_escolhida, '%d/%m/%Y')
                except: pass
                linhas_validas.append((r, val_nf))

        cb_progresso(0.4)

        def chave_ord_padrao(item):
            row, v_nf = item
            return f"{'' if row[0] is None else str(row[0]).strip()}|{'' if row[7] is None else str(row[7]).strip()}|{'' if row[13] is None else str(row[13]).strip()}|{'' if v_nf is None else str(v_nf).strip()}"

        linhas_validas.sort(key=chave_ord_padrao)
        cb_progresso(0.5)

        if ws.max_row > 1: ws.delete_rows(2, ws.max_row)
        for item in linhas_validas: ws.append(item[0][:qtd_colunas_originais])
        
        for row_idx in range(2, len(linhas_validas) + 2):
            ws[f"C{row_idx}"].number_format = 'dd/mm/yyyy'
            for col_letter in ["AA", "AG", "AN", "AQ"]: ws[f"{col_letter}{row_idx}"].number_format = '#,##0.00'
            ws[f"AE{row_idx}"].number_format = '0.00'

        cb_progresso(0.6)

        ws_drive = wb.create_sheet(title="DRIVE")
        colunas_padrao = [1, 3, 59, None, 8, 14, 15, 16, 18, 19, 20, 34, 36, 27, 28, 31, 38, 40, 41, 43, 10, 11, 12, 13, None, None, 54, 57, 60]
        borda_fina = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        fonte_padrao = Font(name='Calibri', size=11)
        fonte_cabecalho = Font(name='Calibri', size=11, bold=True)
        alinhamento_central = Alignment(horizontal='center', vertical='center')
        max_lengths = [0] * len(colunas_padrao)

        for col_drive_idx, col_orig_idx in enumerate(colunas_padrao, start=1):
            celula_cabecalho = ws_drive.cell(row=1, column=col_drive_idx)
            if col_orig_idx is not None:
                val = ws.cell(row=1, column=col_orig_idx).value
                celula_cabecalho.value = val
                if val: max_lengths[col_drive_idx - 1] = max(max_lengths[col_drive_idx - 1], len(str(val)))
            celula_cabecalho.font = fonte_cabecalho
            celula_cabecalho.border = borda_fina
            celula_cabecalho.alignment = alinhamento_central

        chave_anterior = None
        cor_index = -1
        cor_atual = "FFFF00"
        preenchimento = PatternFill(start_color=cor_atual, end_color=cor_atual, fill_type="solid")

        for row_idx, item in enumerate(linhas_validas, start=2):
            r, val_nf = item
            chave = f"{'' if r[0] is None else str(r[0]).strip()}|{'' if r[7] is None else str(r[7]).strip()}|{'' if r[13] is None else str(r[13]).strip()}|{'' if val_nf is None else str(val_nf).strip()}"
            if chave != chave_anterior:
                if chave_anterior is None: cor_atual = "FFFF00"
                else:
                    cor_index = (cor_index + 1) % len(LISTA_CORES)
                    cor_atual = LISTA_CORES[cor_index]
                chave_anterior = chave
                preenchimento = PatternFill(start_color=cor_atual, end_color=cor_atual, fill_type="solid")

            nova_linha = []
            for col_idx_zero, col_orig_idx in enumerate(colunas_padrao):
                val = None
                if col_orig_idx is not None and (col_orig_idx - 1) < len(r):
                    val = r[col_orig_idx - 1]
                nova_linha.append(val)
                if val is not None:
                    tamanho = 10 if isinstance(val, (datetime.date, datetime.datetime)) else len(str(val))
                    if tamanho > max_lengths[col_idx_zero]: max_lengths[col_idx_zero] = tamanho

            ws_drive.append(nova_linha)
            for col_idx, cel in enumerate(ws_drive[row_idx], start=1):
                cel.fill = preenchimento
                cel.font = fonte_padrao
                cel.border = borda_fina
                cel.alignment = alinhamento_central
                if col_idx == 2: cel.number_format = 'dd/mm/yyyy'
            
            if row_idx % 50 == 0: cb_progresso(0.6 + (0.3 * (row_idx / len(linhas_validas))))

        for idx, length in enumerate(max_lengths, start=1):
            if length > 0: ws_drive.column_dimensions[get_column_letter(idx)].width = length + 2

        cb_progresso(0.9)
        wb.save(caminho_final)
        wb.close()
        
        self.criar_dinamica_padrao_excel(caminho_final)
        cb_progresso(1.0)

    def criar_dinamica_padrao_excel(self, arquivo_path):
        caminho_absoluto = os.path.abspath(arquivo_path)
        excel, wb = None, None
        try:
            excel = abrir_excel_silencioso()
            if excel is None: raise Exception("Falha ao abrir Excel no Windows.")
            excel, wb = abrir_workbook_com_retry(excel, caminho_absoluto)
            if wb is None: raise Exception("Falha ao abrir o Workbook.")

            try: ws_dados = wb.Worksheets("Programação Filtrada")
            except: ws_dados = wb.Worksheets(1)

            for sheet in list(wb.Worksheets):
                if sheet.Name == "Dinâmicas Padrão": sheet.Delete()

            ws_dinamica = wb.Worksheets.Add(After=ws_dados)
            ws_dinamica.Name = _normalizar_titulo_planilha("Dinâmicas Padrão")

            ultima_linha = ws_dados.Cells(ws_dados.Rows.Count, 1).End(-4162).Row
            ultima_coluna = ws_dados.Cells(1, ws_dados.Columns.Count).End(-4159).Column
            r1c1_address = f"'{ws_dados.Name}'!R1C1:R{ultima_linha}C{ultima_coluna}"
            pc = wb.PivotCaches().Create(SourceType=1, SourceData=r1c1_address)

            pt1 = pc.CreatePivotTable(TableDestination=ws_dinamica.Range("A4"), TableName="Dinamica1")
            pt1.TableStyle2 = "PivotStyleMedium6"
            pt1.RowAxisLayout(1)

            campo_ciclo = obter_campo_dinamico(pt1, 'Ciclo')
            if campo_ciclo:
                try: campo_ciclo.Orientation = 3
                except: pass

            for i, nome_campo in enumerate(['Data de Programação', 'Loading', 'Frete', 'Razão Social', 'Cidade', 'UF', 'Transportador (Nome)']):
                f = obter_campo_dinamico(pt1, nome_campo)
                if f:
                    try: f.Orientation, f.Position, f.Subtotals = 1, i + 1, [False] * 12
                    except: pass

            for nome, caption in [('Qty', 'Soma de Qty'), ('Cubagem M³', 'Soma de Cubagem M³'), ('Ttl Faturado R$', 'Soma de Ttl Faturado R$')]:
                f = obter_campo_dinamico(pt1, nome)
                if f:
                    try:
                        val = pt1.AddDataField(f)
                        val.Function, val.Caption, val.NumberFormat = -4157, caption, "General"
                    except: pass
            pt1.ColumnGrand, pt1.RowGrand = True, False

            pt2 = pc.CreatePivotTable(TableDestination=ws_dinamica.Range("M4"), TableName="Dinamica2")
            pt2.TableStyle2 = "PivotStyleMedium6"
            pt2.RowAxisLayout(1)

            for i, nome_campo in enumerate(['Loading', 'Pedido', 'Razão Social', 'Código Evap.', 'Cód. Cond.']):
                f = obter_campo_dinamico(pt2, nome_campo)
                if f:
                    try: f.Orientation, f.Position, f.Subtotals = 1, i + 1, [False] * 12
                    except: pass

            f_qty = obter_campo_dinamico(pt2, 'Qty')
            if f_qty:
                try:
                    val3 = pt2.AddDataField(f_qty)
                    val3.Function, val3.Caption, val3.NumberFormat = -4157, 'Total ', "General"
                except: pass
            pt2.ColumnGrand, pt2.RowGrand = True, False
            ws_dinamica.Columns("A:P").AutoFit()
            wb.Save()
        finally:
            if wb:
                try: wb.Close(SaveChanges=False)
                except: pass
            if excel:
                try: excel.Quit()
                except: pass
            try: pythoncom.CoUninitialize()
            except: pass

    def formatar_gmax(self, arquivo_path, data_escolhida, caminho_final, obs_cliente, obs_adic, cb_progresso):
        cb_progresso(0.1)
        wb = openpyxl.load_workbook(arquivo_path)
        ws = wb.active
        qtd_colunas_originais = ws.max_column
        
        cabecalhos = [str(ws.cell(row=1, column=c).value).strip().lower() for c in range(1, ws.max_column + 1)]
        idx_obs_cliente, idx_obs_adc = -1, -1
        
        for i, cab in enumerate(cabecalhos):
            cab_clean = cab.replace(".", "").replace(" ", "").strip()
            if "cliente" in cab_clean and "obs" in cab_clean: idx_obs_cliente = i
            elif "adicionais" in cab_clean and ("obs" in cab_clean or "observa" in cab_clean): idx_obs_adc = i

        if idx_obs_cliente == -1 or idx_obs_adc == -1:
            for i, cab in enumerate(cabecalhos):
                cab_clean = cab.replace(".", "").replace(" ", "").strip()
                if "obs" in cab_clean or "observa" in cab_clean:
                    if i != idx_obs_cliente and i != idx_obs_adc:
                        if idx_obs_cliente == -1: idx_obs_cliente = i
                        elif idx_obs_adc == -1: idx_obs_adc = i

        ws_copia = wb.copy_worksheet(ws)
        ws_copia.title = "Programação Completo"
        ws.title = "Programação Filtrada"

        linhas_raw = list(ws.iter_rows(min_row=2, values_only=True))
        linhas_validas = []
        req_len = max(ws.max_column, 50)
        
        cb_progresso(0.3)

        for r_tuple in linhas_raw:
            r = list(r_tuple)
            if len(r) < req_len: r.extend([None] * (req_len - len(r)))

            v_cliente = r[idx_obs_cliente] if idx_obs_cliente != -1 and idx_obs_cliente < len(r) else None
            v_adc = r[idx_obs_adc] if idx_obs_adc != -1 and idx_obs_adc < len(r) else None
            txt_cliente = str(v_cliente).strip() if v_cliente is not None else ""
            txt_adc = str(v_adc).strip() if v_adc is not None else ""
            
            # =======================================================================
            # CONVERSÃO NUMÉRICA DO GMAX (Garante formato numérico e não string)
            # =======================================================================
            r[0] = converter_numero(r[0])
            r[26] = converter_numero(r[26]) # Qty (Coluna AA)
            
            for i in [30, 31]:              # Preço Unit. (AE) e Ttl. Faturado (AF)
                r[i] = moeda_americana_para_brasileira(r[i])
                
            r[29] = converter_cubagem(r[29]) # Cubagem M³ (Coluna AD)
            # =======================================================================

            # Filtro de datas
            v = r[2]
            d = v.strftime('%d/%m/%Y') if isinstance(v, (datetime.date, datetime.datetime)) else str(v).strip()
            if d == data_escolhida:
                val_nf = txt_cliente if txt_cliente else txt_adc
                linhas_validas.append((r, val_nf))

        cb_progresso(0.5)

        if ws.max_row > 1: ws.delete_rows(2, ws.max_row)
        for item in linhas_validas: ws.append(item[0][:qtd_colunas_originais])
        
        for row_idx in range(2, len(linhas_validas) + 2):
            ws[f"C{row_idx}"].number_format = 'dd/mm/yyyy'
            ws[f"AD{row_idx}"].number_format = '0.00'
            for col_letter in ["AE", "AF"]: 
                ws[f"{col_letter}{row_idx}"].number_format = '#,##0.00'

        cb_progresso(0.6)

        ws_drive = wb.create_sheet(title="DRIVE")
        colunas_gmax = [1, 3, 46, None, 8, 14, 15, 16, 18, 19, 22, 33, 37, 31, 27, 30, 35, 31, 34, None, 10, 11, 12, 13, None, None, 41, 44, 47]
        borda_fina = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        fonte_padrao = Font(name='Calibri', size=11)
        fonte_cabecalho = Font(name='Calibri', size=11, bold=True)
        alinhamento_central = Alignment(horizontal='center', vertical='center')
        
        max_lengths = [0] * len(colunas_gmax)
        for col_drive_idx, col_orig_idx in enumerate(colunas_gmax, start=1):
            celula_cabecalho = ws_drive.cell(row=1, column=col_drive_idx)
            if col_orig_idx is not None:
                val = ws.cell(row=1, column=col_orig_idx).value
                celula_cabecalho.value = val
                if val: max_lengths[col_drive_idx - 1] = max(max_lengths[col_drive_idx - 1], len(str(val)))
            celula_cabecalho.font = fonte_cabecalho
            celula_cabecalho.border = borda_fina
            celula_cabecalho.alignment = alinhamento_central

        chave_anterior = None
        cor_index = -1
        cor_atual = "FFFF00"
        preenchimento = PatternFill(start_color=cor_atual, end_color=cor_atual, fill_type="solid")

        for row_idx, item in enumerate(linhas_validas, start=2):
            r, val_nf = item
            v_a = "" if r[0] is None else str(r[0]).strip()
            v_e = "" if r[7] is None else str(r[7]).strip()
            v_f = "" if r[13] is None else str(r[13]).strip()
            v_ac = "" if r[46] is None else str(r[46]).strip()
            v_nf_str = "" if val_nf is None else str(val_nf).strip()
            
            chave = f"{v_a}|{v_e}|{v_f}|{v_ac}|{v_nf_str}"
            if chave != chave_anterior:
                if chave_anterior is None: cor_atual = "FFFF00"
                else:
                    cor_index = (cor_index + 1) % len(LISTA_CORES)
                    cor_atual = LISTA_CORES[cor_index]
                chave_anterior = chave
                preenchimento = PatternFill(start_color=cor_atual, end_color=cor_atual, fill_type="solid")

            nova_linha = []
            for col_idx_zero, col_orig_idx in enumerate(colunas_gmax):
                val = None
                if col_orig_idx is not None and (col_orig_idx - 1) < len(r): val = r[col_orig_idx - 1]
                nova_linha.append(val)
                if val is not None:
                    tamanho = 10 if isinstance(val, (datetime.date, datetime.datetime)) else len(str(val))
                    if tamanho > max_lengths[col_idx_zero]: max_lengths[col_idx_zero] = tamanho

            ws_drive.append(nova_linha)
            for col_idx, cel in enumerate(ws_drive[row_idx], start=1):
                cel.fill = preenchimento
                cel.font = fonte_padrao
                cel.border = borda_fina
                cel.alignment = alinhamento_central
                if col_idx == 2: cel.number_format = 'dd/mm/yyyy'

        for idx, length in enumerate(max_lengths, start=1):
            if length > 0: ws_drive.column_dimensions[get_column_letter(idx)].width = length + 2

        cb_progresso(0.9)
        wb.save(caminho_final)
        wb.close()
        
        self.criar_dinamica_gmax_excel(caminho_final)
        cb_progresso(1.0)

    def criar_dinamica_gmax_excel(self, arquivo_path):
        caminho_absoluto = os.path.abspath(arquivo_path)
        excel, wb = None, None
        try:
            excel = abrir_excel_silencioso()
            if excel is None: raise Exception("Falha ao abrir Excel no Windows.")
            excel, wb = abrir_workbook_com_retry(excel, caminho_absoluto)
            if wb is None: raise Exception("Falha ao abrir o Workbook.")

            try: ws_dados = wb.Worksheets("Programação Filtrada")
            except: ws_dados = wb.Worksheets(1)

            for sheet in list(wb.Worksheets):
                if sheet.Name == "Dinâmica G-MAX": sheet.Delete()

            ws_dinamica = wb.Worksheets.Add(After=ws_dados)
            ws_dinamica.Name = _normalizar_titulo_planilha("Dinâmica G-MAX")

            ultima_linha = ws_dados.Cells(ws_dados.Rows.Count, 1).End(-4162).Row
            ultima_coluna = ws_dados.Cells(1, ws_dados.Columns.Count).End(-4159).Column
            r1c1_address = f"'{ws_dados.Name}'!R1C1:R{ultima_linha}C{ultima_coluna}"
            pc = wb.PivotCaches().Create(SourceType=1, SourceData=r1c1_address)

            pt1 = pc.CreatePivotTable(TableDestination=ws_dinamica.Range("A4"), TableName="DinamicaGMAX1")
            pt1.TableStyle2 = "PivotStyleMedium6"
            pt1.RowAxisLayout(1)

            f_unidade = obter_campo_dinamico(pt1, 'Unidade')
            if f_unidade:
                try: f_unidade.Orientation = 3
                except: pass

            for i, nome_campo in enumerate(['Data de Programação', 'Loading', 'Frete', 'Razão Social', 'Cidade', 'UF', 'Transportador']):
                f = obter_campo_dinamico(pt1, nome_campo)
                if f:
                    try: f.Orientation, f.Position, f.Subtotals = 1, i + 1, [False] * 12
                    except: pass

            for nome, caption in [('Qty', 'Soma de Qty'), ('Cubagem M³', 'Soma de Cubagem M³'), ('Ttl. Faturado (R $)', 'Soma de Ttl. Faturado (R$ )')]:
                f = obter_campo_dinamico(pt1, nome)
                if f:
                    try:
                        val = pt1.AddDataField(f)
                        val.Function, val.Caption, val.NumberFormat = -4157, caption, "General"
                    except: pass
            pt1.ColumnGrand, pt1.RowGrand = True, False

            pt2 = pc.CreatePivotTable(TableDestination=ws_dinamica.Range("M4"), TableName="DinamicaGMAX2")
            pt2.TableStyle2 = "PivotStyleMedium6"
            pt2.RowAxisLayout(1)

            for i, nome_campo in enumerate(['Loading', 'Pedido', 'Razão Social', 'Código Unidade', 'Unidade']):
                f = obter_campo_dinamico(pt2, nome_campo)
                if f:
                    try: f.Orientation, f.Position, f.Subtotals = 1, i + 1, [False] * 12
                    except: pass

            f_qty = obter_campo_dinamico(pt2, 'Qty')
            if f_qty:
                try:
                    val3 = pt2.AddDataField(f_qty)
                    val3.Function, val3.Caption, val3.NumberFormat = -4157, 'Soma Qty ', "General"
                except: pass
            pt2.ColumnGrand, pt2.RowGrand = True, False
            ws_dinamica.Columns("A:R").AutoFit()
            wb.Save()
        finally:
            if wb:
                try: wb.Close(SaveChanges=False)
                except: pass
            if excel:
                try: excel.Quit()
                except: pass
            try: pythoncom.CoUninitialize()
            except: pass